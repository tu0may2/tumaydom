from openpyxl import load_workbook

from fitcoach.db import Database
from fitcoach.export import _format_exercises, export_workbook

SHEETS = ["Обзор", "Недели", "Прогресс", "Тренировки", "Сон", "Замеры"]


def _seed(tmp_path) -> Database:
    db = Database(tmp_path / "t.db")
    db.update_profile(1, {"goal": "сила"}, chat_id=1)
    for day, weight in (("2026-08-11", 90), ("2026-08-18", 100)):
        db.add_workout(1, {"date": day, "kind": "strength", "title": "Ноги", "rpe": 8,
                           "exercises": [{"name": "присед",
                                          "sets": [{"weight_kg": weight, "reps": 5}] * 4}]})
        db.add_sleep(1, {"date": day, "total_min": 420, "score": 75, "hrv_ms": 60})
        db.add_metric(1, "weight_kg", 82.5, day)
    return db


def test_export_creates_all_sheets(tmp_path):
    db = _seed(tmp_path)
    path = export_workbook(db, 1, tmp_path / "out.xlsx")

    assert path.exists() and path.stat().st_size > 0
    workbook = load_workbook(path)
    assert workbook.sheetnames == SHEETS


def test_export_fills_progress_row(tmp_path):
    db = _seed(tmp_path)
    workbook = load_workbook(export_workbook(db, 1, tmp_path / "out.xlsx"))
    sheet = workbook["Прогресс"]
    assert sheet.cell(row=1, column=1).value == "Упражнение"
    assert sheet.cell(row=2, column=1).value == "присед"
    assert sheet.cell(row=2, column=8).value == "100×5"


def test_export_without_data_still_writes_file(tmp_path):
    db = Database(tmp_path / "empty.db")
    path = export_workbook(db, 42, tmp_path / "empty.xlsx")
    workbook = load_workbook(path)
    assert workbook["Тренировки"].max_row == 1  # только заголовок


def test_format_exercises_renders_sets_and_holds():
    text = _format_exercises([
        {"name": "жим", "sets": [{"weight_kg": 80, "reps": 5}, {"weight_kg": 82.5, "reps": 3}]},
        {"name": "планка", "sets": [{"duration_sec": 60}]},
    ])
    assert text == "жим: 80×5, 82.5×3; планка: 60 сек"
