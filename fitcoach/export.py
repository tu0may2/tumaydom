"""Выгрузка статистики в Excel: цифры + разбор причин прогресса и провалов."""

from __future__ import annotations

import logging
from datetime import date
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analysis.progress import exercise_progress, flags, weekly_factors
from .analysis.stats import metric_trend, sleep_trend
from .db import Database
from .llm import LLM, EFFORT_ANALYSIS
from .prompts import coach_system

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=13, bold=True)

INSIGHT_PROMPT = """Ниже — статистика спортсмена за период. Напиши разбор для листа
Excel: почему результат такой, какой есть.

Структура ответа, каждый пункт с новой строки, без нумерации и без markdown:
- «Что работает:» 2–3 строки — какие решения дали прогресс, с опорой на цифры.
- «Что тормозит:» 2–3 строки — где просадка и чем она объясняется в данных
  (объём, сон, HRV, пропуски, вес). Если данных не хватает для вывода — скажи это
  прямо, а не придумывай причину.
- «Что менять:» 2–3 конкретных действия на ближайшие две недели.

Не пересказывай таблицы, объясняй связи."""


def export_workbook(db: Database, user_id: int, path: str | Path, *,
                    days: int = 180, llm: LLM | None = None) -> Path:
    """Собрать .xlsx. С `llm` добавляется лист с текстовым разбором."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    workouts = db.recent_workouts(user_id, days=days)
    nights = db.recent_sleep(user_id, days=days)
    metrics = db.recent_metrics(user_id, days=days)
    factors = weekly_factors(workouts, nights, metrics, weeks=52)
    progress = exercise_progress(workouts)
    warnings = flags(factors)

    workbook = Workbook()
    workbook.remove(workbook.active)

    _sheet_overview(workbook, db, user_id, workouts, nights, metrics, factors, warnings)
    _sheet_weeks(workbook, factors)
    _sheet_exercises(workbook, progress)
    _sheet_workouts(workbook, workouts)
    _sheet_sleep(workbook, nights)
    _sheet_metrics(workbook, metrics)

    if llm is not None:
        try:
            _sheet_insight(workbook, llm, factors, progress, warnings)
        except Exception:  # разбор — приятное дополнение, файл важнее
            log.exception("Не удалось получить текстовый разбор для выгрузки")

    workbook.save(path)
    return path


# ------------------------------------------------------------------- листы


def _sheet_overview(workbook: Workbook, db: Database, user_id: int,
                    workouts: list[dict[str, Any]], nights: list[dict[str, Any]],
                    metrics: list[dict[str, Any]], factors: list[dict[str, Any]],
                    warnings: list[str]) -> None:
    sheet = workbook.create_sheet("Обзор")
    sheet["A1"] = "Сводка FitCoach"
    sheet["A1"].font = TITLE_FONT

    profile = db.get_profile(user_id)
    trends = metric_trend(metrics)
    sleep = sleep_trend(nights)
    total_tonnage = sum(w.get("tonnage_kg") or 0 for w in workouts)
    total_distance = sum(w.get("distance_km") or 0 for w in workouts)

    rows: list[tuple[str, Any]] = [
        ("Выгрузка от", date.today().isoformat()),
        ("Цель", profile.get("goal", "не задана")),
        ("Тренировок за период", len(workouts)),
        ("Силовых", sum(1 for w in workouts if w.get("kind") == "strength")),
        ("Кардио", sum(1 for w in workouts if w.get("kind") == "cardio")),
        ("Суммарный тоннаж, кг", round(total_tonnage, 1)),
        ("Суммарная дистанция, км", round(total_distance, 2)),
        ("Ночей с данными сна", len(nights)),
    ]
    if sleep.get("avg_7d", {}).get("total_min"):
        rows.append(("Средний сон за 7 дней, ч",
                     round(sleep["avg_7d"]["total_min"] / 60, 2)))
    for name, trend in trends.items():
        rows.append((f"{name}: последнее", trend["last"]))
        rows.append((f"{name}: изменение за период", trend["change"]))

    line = 3
    for label, value in rows:
        sheet.cell(row=line, column=1, value=label).font = Font(bold=True)
        sheet.cell(row=line, column=2, value=value)
        line += 1

    line += 1
    sheet.cell(row=line, column=1, value="На что посмотреть").font = TITLE_FONT
    line += 1
    for note in warnings or ["Явных тревожных сигналов в данных нет."]:
        cell = sheet.cell(row=line, column=1, value=note)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.merge_cells(start_row=line, start_column=1, end_row=line, end_column=6)
        line += 1

    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 22


def _sheet_weeks(workbook: Workbook, factors: list[dict[str, Any]]) -> None:
    headers = [
        ("week_start", "Неделя с"),
        ("sessions", "Сессий"),
        ("tonnage_kg", "Тоннаж, кг"),
        ("tonnage_change_pct", "Тоннаж, % к пред."),
        ("distance_km", "Дистанция, км"),
        ("duration_min", "Время, мин"),
        ("avg_rpe", "Средний RPE"),
        ("avg_sleep_min", "Сон, мин"),
        ("avg_sleep_score", "Sleep score"),
        ("avg_hrv_ms", "HRV, мс"),
        ("avg_resting_hr", "Пульс покоя"),
        ("avg_weight_kg", "Вес, кг"),
    ]
    ordered = sorted(factors, key=lambda r: r["week_start"])
    sheet = _table(workbook, "Недели", headers, ordered)

    if len(ordered) >= 2:
        _add_chart(sheet, len(ordered), value_col=3, title="Тоннаж по неделям", anchor="N2")
        _add_chart(sheet, len(ordered), value_col=8, title="Сон по неделям, мин", anchor="N20")


def _sheet_exercises(workbook: Workbook, progress: list[dict[str, Any]]) -> None:
    headers = [
        ("exercise", "Упражнение"),
        ("sessions", "Сессий"),
        ("first_date", "Первая"),
        ("last_date", "Последняя"),
        ("first_e1rm", "1ПМ на старте"),
        ("last_e1rm", "1ПМ сейчас"),
        ("best_e1rm", "Лучший 1ПМ"),
        ("best_set", "Лучший подход"),
        ("change_kg", "Δ кг"),
        ("change_pct", "Δ %"),
        ("last_volume_kg", "Объём посл., кг"),
    ]
    _table(workbook, "Прогресс", headers, progress)


def _sheet_workouts(workbook: Workbook, workouts: list[dict[str, Any]]) -> None:
    headers = [
        ("date", "Дата"),
        ("kind", "Тип"),
        ("title", "Название"),
        ("duration_min", "Мин"),
        ("rpe", "RPE"),
        ("tonnage_kg", "Тоннаж, кг"),
        ("distance_km", "Км"),
        ("avg_hr", "Ср. пульс"),
        ("max_hr", "Макс. пульс"),
        ("exercises_text", "Упражнения"),
        ("notes", "Заметки"),
        ("source", "Источник"),
    ]
    rows = [{**w, "exercises_text": _format_exercises(w.get("exercises") or [])}
            for w in workouts]
    _table(workbook, "Тренировки", headers, rows)


def _sheet_sleep(workbook: Workbook, nights: list[dict[str, Any]]) -> None:
    headers = [
        ("date", "Дата"),
        ("total_min", "Всего, мин"),
        ("deep_min", "Глубокий"),
        ("rem_min", "REM"),
        ("light_min", "Лёгкий"),
        ("awake_min", "Бодрствование"),
        ("score", "Score"),
        ("resting_hr", "Пульс покоя"),
        ("hrv_ms", "HRV, мс"),
        ("body_battery", "Body battery"),
        ("stress_avg", "Стресс"),
        ("spo2", "SpO2"),
    ]
    _table(workbook, "Сон", headers, sorted(nights, key=lambda n: n["date"]))


def _sheet_metrics(workbook: Workbook, metrics: list[dict[str, Any]]) -> None:
    headers = [("date", "Дата"), ("name", "Показатель"), ("value", "Значение")]
    _table(workbook, "Замеры", headers, sorted(metrics, key=lambda m: m["date"]))


def _sheet_insight(workbook: Workbook, llm: LLM, factors: list[dict[str, Any]],
                   progress: list[dict[str, Any]], warnings: list[str]) -> None:
    import json

    payload = {
        "weeks": factors[:12],
        "exercise_progress": progress[:20],
        "flags": warnings,
    }
    text = llm.text(
        f"{INSIGHT_PROMPT}\n\nДанные:\n{json.dumps(payload, ensure_ascii=False, indent=2)}",
        system=coach_system(),
        effort=EFFORT_ANALYSIS,
        max_tokens=4000,
    )

    sheet = workbook.create_sheet("Разбор", 1)
    sheet["A1"] = "Разбор: что работает, что тормозит"
    sheet["A1"].font = TITLE_FONT
    sheet.column_dimensions["A"].width = 110

    line = 3
    for paragraph in text.split("\n"):
        cell = sheet.cell(row=line, column=1, value=paragraph)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if paragraph.strip().endswith(":"):
            cell.font = Font(bold=True)
        line += 1


# ------------------------------------------------------------------ helpers


def _table(workbook: Workbook, title: str, headers: Sequence[tuple[str, str]],
           rows: list[dict[str, Any]]):
    sheet = workbook.create_sheet(title)

    for column, (_, label) in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for line, row in enumerate(rows, start=2):
        for column, (key, _) in enumerate(headers, start=1):
            sheet.cell(row=line, column=column, value=row.get(key))

    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        )
    _autosize(sheet, len(headers))
    return sheet


def _autosize(sheet, columns: int, limit: int = 46) -> None:
    for column in range(1, columns + 1):
        longest = max(
            (len(str(cell.value)) for cell in sheet[get_column_letter(column)]
             if cell.value is not None),
            default=8,
        )
        sheet.column_dimensions[get_column_letter(column)].width = min(longest + 2, limit)


def _add_chart(sheet, row_count: int, *, value_col: int, title: str, anchor: str) -> None:
    chart = LineChart()
    chart.title = title
    chart.height, chart.width = 8, 16
    data = Reference(sheet, min_col=value_col, min_row=1, max_row=row_count + 1)
    categories = Reference(sheet, min_col=1, min_row=2, max_row=row_count + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    sheet.add_chart(chart, anchor)


def _format_exercises(exercises: list[dict[str, Any]]) -> str:
    parts = []
    for exercise in exercises:
        sets = exercise.get("sets") or []
        rendered = ", ".join(
            f"{s['weight_kg']:g}×{s['reps']}" if s.get("weight_kg") and s.get("reps")
            else (f"{s['reps']} повт." if s.get("reps") else f"{s.get('duration_sec', 0):g} сек")
            for s in sets
        )
        parts.append(f"{exercise.get('name')}: {rendered}" if rendered
                     else str(exercise.get("name")))
    return "; ".join(parts)
